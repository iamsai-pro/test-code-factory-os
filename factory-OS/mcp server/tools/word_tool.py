from docx import Document
from .common import safe_path

def register(mcp):
    @mcp.tool()
    def word_create(path: str, title: str, body: str) -> str:
        """Create a DOCX document with a title and body paragraphs separated by blank lines."""
        p = safe_path(path)
        if p.suffix.lower() != ".docx": raise ValueError("Output path must end in .docx")
        p.parent.mkdir(parents=True, exist_ok=True)
        doc = Document(); doc.add_heading(title, 0)
        for para in body.split("\n\n"): doc.add_paragraph(para)
        doc.save(p)
        return f"Created {p.name}"

    @mcp.tool()
    def word_read(path: str, max_chars: int = 20000) -> str:
        """Extract paragraph text from a DOCX file in the workspace."""
        p = safe_path(path); doc = Document(p)
        text = "\n".join(para.text for para in doc.paragraphs)
        return text[:max(1, min(max_chars, 100000))]
