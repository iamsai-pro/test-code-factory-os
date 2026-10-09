import json
from pptx import Presentation
from pptx.util import Inches, Pt
from .common import safe_path

def register(mcp):
    @mcp.tool()
    def powerpoint_create(path: str, slides_json: str) -> str:
        """Create a PPTX from JSON: [{\"title\":\"...\",\"bullets\":[\"...\"]}, ...]."""
        slides = json.loads(slides_json)
        if not isinstance(slides, list) or not slides: raise ValueError("slides_json must be a non-empty JSON array.")
        p = safe_path(path)
        if p.suffix.lower() != ".pptx": raise ValueError("Output path must end in .pptx")
        p.parent.mkdir(parents=True, exist_ok=True)
        prs = Presentation()
        for item in slides:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = str(item.get("title", "Untitled"))
            body = slide.placeholders[1].text_frame
            body.clear()
            for i, bullet in enumerate(item.get("bullets", [])):
                para = body.paragraphs[0] if i == 0 else body.add_paragraph()
                para.text = str(bullet); para.level = 0; para.font.size = Pt(22)
        prs.save(p)
        return f"Created {p.name} with {len(slides)} slides"
