from openpyxl import Workbook, load_workbook
from .common import safe_path

def register(mcp):
    @mcp.tool()
    def excel_create(path: str, sheet_name: str = "Sheet1", headers_csv: str = "") -> str:
        """Create an XLSX workbook. Optional headers are comma-separated."""
        p = safe_path(path)
        if p.suffix.lower() != ".xlsx": raise ValueError("Output path must end in .xlsx")
        p.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook(); ws = wb.active; ws.title = sheet_name[:31] or "Sheet1"
        if headers_csv.strip(): ws.append([x.strip() for x in headers_csv.split(",")])
        wb.save(p)
        return f"Created {p.name}"

    @mcp.tool()
    def excel_read(path: str, sheet_name: str = "", max_rows: int = 100) -> str:
        """Read worksheet values from an XLSX workbook."""
        p = safe_path(path); max_rows = max(1, min(max_rows, 500))
        wb = load_workbook(p, read_only=True, data_only=True)
        ws = wb[sheet_name] if sheet_name else wb.active
        rows = list(ws.iter_rows(values_only=True, max_row=max_rows))
        return "\n".join(" | ".join("" if v is None else str(v) for v in row) for row in rows)

    @mcp.tool()
    def excel_append_rows(path: str, rows_json: str, sheet_name: str = "") -> str:
        """Append rows from a JSON array of arrays, e.g. [[\"A\",1],[\"B\",2]]."""
        import json
        rows = json.loads(rows_json)
        if not isinstance(rows, list) or any(not isinstance(row, list) for row in rows):
            raise ValueError("rows_json must be a JSON array of arrays.")
        p = safe_path(path); wb = load_workbook(p)
        ws = wb[sheet_name] if sheet_name else wb.active
        for row in rows: ws.append(row)
        wb.save(p)
        return f"Appended {len(rows)} rows to {p.name}"
